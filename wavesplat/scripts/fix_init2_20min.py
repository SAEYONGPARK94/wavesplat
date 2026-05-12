"""
fix_init2_20min.py

docs/labmeeting_20260513_박세용_20min.pptx 의 slide 9 (Init 2/2) 를
labmeeting_2026_05_13.pptx 와 동일한 새 구조로 수정.

수정 사항:
1. 순서 재정렬 — Beamforming → Adam (Loss 정의) → Coarse-to-fine N schedule
2. Loss function 식으로 명시 (L_data, L_L1, L_smooth)
3. Coarse-to-fine 의 *Loss* 가 Adam 의 fit Loss 임을 명확히
4. Stage advance 조건: (a) Loss<thr OR (b) max_iter 모두 표시
"""
import os
import shutil
from pptx import Presentation
from pptx.util import Pt, Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PATH = os.path.join(ROOT, 'docs', 'labmeeting_20260513_박세용_20min.pptx')

# Backup
BACKUP = PATH.replace('.pptx', '_backup.pptx')


# New slide 9 body content (사용자 compact tone, Adam 은 간단히)
SLIDE_9_INIT2 = [
    (0, 'Init (2/2) — Beamforming → Adam (with coarse-to-fine N schedule)'),
    (1, '① Beamforming score (초기 (k_x, k_y) 글로벌 탐색)'),
    (2, 'S(k) = | Σ_i η_obs(x_i, t_i) · exp(−j(k·x_i − ω(k)·t_i)) |²'),
    (2, 'i: init window 의 모든 점 (frame × 점 수 = 수천 개)'),
    (2, 'ω(k) 는 dispersion 으로 자동 — 2D search (k_x, k_y) 만'),
    (2, 'Search range: k_x, k_y ∈ [−4, 4], 81×81 = 6,561 후보'),
    (2, '→ top-N peak 선택 (Adam 초기값으로 전달)'),
    (1, '② Adam refinement (multi-frame fitting)'),
    (2, 'Loss = data fit + L1 sparsity + spectral smoothness 의 3 항 (수식은 노트 참조)'),
    (2, 'Multi-frame init window 전체에 fitting'),
    (2, '파라미터 (A, k_x, k_y, φ) 동시 update, ω 는 dispersion 으로 자동'),
    (2, 'k 는 lr 의 절반 (0.005) 으로 천천히 — 변화 신중'),
    (1, '③ Coarse-to-fine N schedule (Adam stage 진행 방식)'),
    (2, 'N: 3 → 8 → 15 (성분 수 점진 증가, 큰 성분 먼저)'),
    (2, 'Stage advance 조건 — 둘 중 하나만 만족 시:'),
    (2, '   (a) Loss < threshold (stage 1: 0.1, stage 2: 0.05, stage 3: None)'),
    (2, '   (b) Max iter 도달 (stage 1·2: 200, stage 3: 500 — 총 ≤ 900)'),
    (2, '빨리 수렴된 stage 는 즉시 advance, 못 수렴이면 iter cap 으로 강제 advance'),
]


def find_body_placeholder(slide):
    """Find body placeholder (large text frame)."""
    candidates = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        t = shape.top / 914400 if shape.top else 0
        h = shape.height / 914400 if shape.height else 0
        if t > 0.5 and h > 4:
            candidates.append((shape, h))
    if candidates:
        candidates.sort(key=lambda x: -x[1])
        return candidates[0][0]
    return None


def fill_body(text_frame, items, font_sizes=(13, 11, 10), space_after=(2, 1, 0)):
    """Fill body with (level, text) tuples."""
    text_frame.word_wrap = True
    p0 = text_frame.paragraphs[0]
    p0.text = ''
    while len(text_frame.paragraphs) > 1:
        text_frame._txBody.remove(text_frame.paragraphs[-1]._p)
    for i, (level, text) in enumerate(items):
        if i == 0:
            p = p0
        else:
            p = text_frame.add_paragraph()
        p.text = text
        p.level = level
        sz = font_sizes[min(level, len(font_sizes) - 1)]
        sp = space_after[min(level, len(space_after) - 1)]
        p.space_after = Pt(sp)
        for r in p.runs:
            r.font.size = Pt(sz)
            if level == 0:
                r.font.bold = True


NOTES_INIT2 = (
    "Init 은 *Beamforming → Adam* 두 단계로 작동합니다. Beamforming 으로 초기 (k_x, k_y) "
    "후보를 글로벌 탐색하고, 그 결과를 Adam 의 초기값으로 넘겨서 정밀화하는 흐름이에요. "
    "Adam 안에서는 *coarse-to-fine N schedule* 로 성분 수를 점진적으로 늘려갑니다.\n\n"
    "(1) **Beamforming score** — 코러스 효과로 설명드리면, 맞는 wavenumber 면 모든 점의 "
    "phase 가 일치해서 합이 N 배로 커지고, 틀린 wavenumber 면 random cancel 로 √N 으로만 "
    "커집니다. 81×81 격자 후보 중 top-N peak 를 Adam 초기값으로 넘깁니다.\n\n"
    "(2) **Adam refinement — Loss 최소화** — Beamforming 결과는 격자 위 peak 라 거칠어요. "
    "그래서 Adam optimizer 로 multi-frame init window 전체에 다시 fitting 합니다. *Loss 함수* "
    "는 3 항으로 구성됩니다:\n\n"
    "  Loss(θ) = L_data(θ) + λ_L1·L_L1(θ) + λ_smooth·L_smooth(θ)\n\n"
    "  - L_data(θ) = (1/N_obs) · Σ_i c_i · [η_model(x_i, t_i; θ) − η_obs(x_i, t_i)]²\n"
    "    → 측정 점과 model 예측의 confidence-weighted MSE. c_i 는 가까운 점일수록 ↑\n"
    "  - L_L1(θ) = Σ_n |A_n|  (계수 λ_L1 = 0.001)\n"
    "    → amplitude sparsity 유도, 의미 없는 작은 성분 살아남지 않게 함\n"
    "  - L_smooth(θ) = Σ_n (A_↑[n+1] − A_↑[n])²  (계수 λ_smooth = 0.0001)\n"
    "    → |k| 오름차순 정렬 후 인접 amplitude 차분 제곱, spectrum 부드러움 유도\n\n"
    "여기서 η_model(x, y, t; θ) = Σ_n A_n · cos(k_xn·x + k_yn·y − ω_n·t + φ_n) 이고, "
    "파라미터 θ = (A, k_x, k_y, φ), ω 는 dispersion 으로 자동 (별도 자유도 아님). Adam 이 "
    "이 Loss 를 최소화하는 방향으로 update 하는데, k 는 변화에 신중해야 해서 lr 을 절반 "
    "(0.005) 으로 줍니다. (A, φ) 는 default 0.01.\n\n"
    "(3) **Coarse-to-fine N schedule** — N 을 3 → 8 → 15 로 점진 증가. *큰 성분 먼저* "
    "찾고 그 다음 작은 성분을 추가합니다. *Stage advance 는 두 조건 중 하나* 만족 시 — "
    "(a) Loss (위 ②에서 정의한 fit Loss) < threshold 또는 (b) max iter 도달. 여기서의 "
    "Loss 는 beamforming S(k) 와 *다른 함수* — Adam fit Loss 임을 강조해 주세요."
)


def main():
    # Backup
    if not os.path.exists(BACKUP):
        shutil.copy(PATH, BACKUP)
        print(f'[Backup] {BACKUP}')

    prs = Presentation(PATH)
    print(f'Total slides: {len(prs.slides)}')

    # Slide 9 (0-indexed = 8): Init (2/2)
    slide = prs.slides[8]
    body = find_body_placeholder(slide)
    if body is None:
        print('[ERR] Body placeholder not found on slide 9')
        return
    print('  Body placeholder found, updating content...')
    fill_body(body.text_frame, SLIDE_9_INIT2,
              font_sizes=(12, 11, 9), space_after=(2, 1, 0))

    # Update presenter notes
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = NOTES_INIT2
    print('  Presenter notes updated')

    prs.save(PATH)
    print(f'\nSaved: {PATH}')


if __name__ == '__main__':
    main()
